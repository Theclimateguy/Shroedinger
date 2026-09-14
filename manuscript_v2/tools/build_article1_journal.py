#!/usr/bin/env python3
"""Build the journal-format manuscript of article 1 (Izvestiya RAN, Ser. Geogr.).

Input : manuscript_v2/article1_v7_src.tex   -- text with \\cite{key} placeholders
        manuscript_v2/references_article1.bib
Output: manuscript_v2/article1_v7.tex        -- author-year citations in plain text,
                                                unnumbered GOST list + Latin References
                                                (journal rules of 17.11.2025)
        manuscript_v2/article1_v7.docx       -- via pandoc, then Times 12 / 1.5 / margins
        manuscript_v2/article1_v7.pdf        -- via latexmk -xelatex (if available)

Usage: python manuscript_v2/tools/build_article1_journal.py [--no-pdf] [--no-docx]
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
MAN = HERE.parent
SRC = MAN / "article1_v7_src.tex"
BIB = MAN / "references_article1.bib"
OUT_TEX = MAN / "article1_v7.tex"
OUT_DOCX = MAN / "article1_v7.docx"

# ----------------------------------------------------------------------------- bib
def parse_bib(path: Path) -> dict[str, dict]:
    s = path.read_text(encoding="utf-8")
    entries: dict[str, dict] = {}
    i = 0
    while True:
        m = re.compile(r"@(\w+)\s*\{\s*([^,\s]+)\s*,").search(s, i)
        if not m:
            break
        etype, key = m.group(1).lower(), m.group(2)
        j = m.end()
        depth, start = 1, m.start()
        while depth and j < len(s):
            if s[j] == "{":
                depth += 1
            elif s[j] == "}":
                depth -= 1
            j += 1
        body = s[m.end():j - 1]
        fields: dict[str, str] = {"_type": etype}
        k = 0
        while k < len(body):
            fm = re.compile(r"\s*(\w+)\s*=\s*").match(body, k)
            if not fm:
                k += 1
                continue
            name = fm.group(1).lower()
            k = fm.end()
            if body[k] == "{":
                d, q = 1, k + 1
                while d:
                    if body[q] == "{":
                        d += 1
                    elif body[q] == "}":
                        d -= 1
                    q += 1
                val = body[k + 1:q - 1]
                k = q
            elif body[k] == '"':
                q = body.index('"', k + 1)
                val, k = body[k + 1:q], q + 1
            else:
                q = body.find(",", k)
                q = len(body) if q < 0 else q
                val, k = body[k:q].strip(), q
            fields[name] = " ".join(val.split())
            c = body.find(",", k)
            k = len(body) if c < 0 else c + 1
        entries[key] = fields
        i = j
    return entries


LATEX_ACCENTS = {r"{\'e}": "é", r"{\'a}": "á", r"{\'i}": "í", r"{\'o}": "ó", r"{\'u}": "ú",
                 r"\'e": "é", r"\'a": "á", r"\'i": "í", r"{\v{c}}": "č", r"\v{c}": "č",
                 r"{\c{c}}": "ç", r"{\~n}": "ñ", r"\~n": "ñ"}


def clean(s: str) -> str:
    for a, b in LATEX_ACCENTS.items():
        s = s.replace(a, b)
    s = s.replace("--", "–").replace("~", " ")
    s = re.sub(r"[{}]", "", s)
    return s


def split_authors(field: str) -> tuple[list[tuple[str, str]], bool]:
    """-> ([(surname, initials)], has_et_al)"""
    parts = [p.strip() for p in re.split(r"\s+and\s+", field)]
    out, etal = [], False
    for p in parts:
        if p.lower() == "others":
            etal = True
            continue
        if "," in p:
            last, first = [x.strip() for x in p.split(",", 1)]
        else:
            toks = p.split()
            last, first = toks[-1], " ".join(toks[:-1])
        last, first = clean(last), clean(first)
        inits = "".join(t[0] + "." for t in re.split(r"[\s.]+", first) if t)
        inits = re.sub(r"(\w)\.-?(\w)\.", lambda m: m.group(1) + "." + m.group(2) + ".", inits)
        out.append((last, inits))
    return out, etal


def is_cyrillic(s: str) -> bool:
    return bool(re.search("[А-Яа-яЁё]", s))


# ------------------------------------------------------------- transliteration (journal)
TR = {"а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z",
      "и": "i", "й": "i", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r",
      "с": "s", "т": "t", "у": "u", "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh",
      "щ": "shch", "ъ": "’’", "ы": "y", "ь": "’", "э": "e", "ю": "yu", "я": "ya"}


def translit(s: str) -> str:
    out = []
    for ch in s:
        lo = ch.lower()
        if lo in TR:
            t = TR[lo]
            out.append(t.capitalize() if ch.isupper() else t)
        else:
            out.append(ch)
    return "".join(out)


def translit_initials(inits: str) -> str:
    # "Ю.Г." -> "Yu.G."
    return "".join(translit(c) if c != "." else "." for c in inits)


# ------------------------------------------------------------------- journal abbreviations
JOURNAL_ABBR = {
    "Известия РАН. Серия географическая": ("Изв. РАН. Сер. геогр.", "Izv. Akad. Nauk, Ser. Geogr."),
    "Quarterly Journal of the Royal Meteorological Society": ("Q. J. R. Meteorol. Soc.",) * 2,
    "Journal of Climate": ("J. Clim.",) * 2,
    "Journal of Geophysical Research: Atmospheres": ("J. Geophys. Res. Atmos.",) * 2,
    "Journal of Fluid Mechanics": ("J. Fluid Mech.",) * 2,
    "Communications on Pure and Applied Mathematics": ("Commun. Pure Appl. Math.",) * 2,
    "IEEE Transactions on Pattern Analysis and Machine Intelligence": ("IEEE Trans. Pattern Anal. Mach. Intell.",) * 2,
    "PNAS Nexus": ("PNAS Nexus",) * 2,
    "Climate Dynamics": ("Clim. Dyn.",) * 2,
    "Monthly Weather Review": ("Mon. Weather Rev.",) * 2,
    "Journal of the Atmospheric Sciences": ("J. Atmos. Sci.",) * 2,
    "Tellus": ("Tellus",) * 2,
    "Journal of Advances in Modeling Earth Systems": ("J. Adv. Model. Earth Syst.",) * 2,
    "Journal of Physical Oceanography": ("J. Phys. Oceanogr.",) * 2,
    "Annual Review of Fluid Mechanics": ("Annu. Rev. Fluid Mech.",) * 2,
    "Nature Physics": ("Nat. Phys.",) * 2,
    "Geophysical Research Letters": ("Geophys. Res. Lett.",) * 2,
    "Geoscientific Model Development": ("Geosci. Model Dev.",) * 2,
    "Science": ("Science",) * 2,
    "Nature": ("Nature",) * 2,
    "NeuroReport": ("NeuroReport",) * 2,
    "Revue géographique des Pyrénées et du Sud-Ouest": ("Rev. Géogr. Pyrénées Sud-Ouest",) * 2,
    "Proceedings of the American Philosophical Society": ("Proc. Am. Philos. Soc.",) * 2,
    "Ecology": ("Ecology",) * 2,
    "The Quarterly Review of Biology": ("Q. Rev. Biol.",) * 2,
    "Atmospheric Research": ("Atmos. Res.",) * 2,
}

# Sentence-case titles where the bib carries Title Case; English titles of Russian sources.
TITLE_OVERRIDE = {
    "Simon1962": "The architecture of complexity",
    "Levin1992": "The problem of pattern and scale in ecology: The Robert H. MacArthur Award Lecture",
    "WuLoucks1995": "From balance of nature to hierarchical patch dynamics: A paradigm shift in ecology",
    "NastromGage1985": "A climatology of atmospheric wavenumber spectra of wind and temperature observed by commercial aircraft",
    "Skinner2025": "Characterizing ocean flows with the scattering transform",
    "Canolty2006": "High gamma power is phase-locked to theta oscillations in human neocortex",
    "LovejoySchertzer2013": "The Weather and Climate: Emergent Laws and Multifractal Cascades",
    "Bertrand1968": "Paysage et géographie physique globale. Esquisse méthodologique",
    "Haarsma2016": "High Resolution Model Intercomparison Project (HighResMIP v1.0) for CMIP6",
}
EN_TITLE = {  # English titles / translations for Cyrillic sources (References list)
    "Sochava1978": "Introduction to the Theory of Geosystems",
    "Puzachenko1986": "Spatio-temporal hierarchy of geosystems from the standpoint of oscillation theory",
    "Puzachenko2004": "Mathematical Methods in Ecological and Geographical Research",
    "Puzachenko2010": "Invariants of a dynamic geosystem",
    "Krenke2019": "Spatial organization of regional mesoclimate",
}
EN_BOOKTITLE = {
    "Puzachenko1986": ("Voprosy geografii. Sb. 127: Modelirovanie geosistem",
                       "Problems of Geography. Vol. 127: Modelling of Geosystems"),
}
CITY_RU = {"Москва": "М.", "Новосибирск": "Новосибирск", "Cambridge": "Cambridge"}
CITY_EN = {"Москва": "Moscow", "Новосибирск": "Novosibirsk", "Cambridge": "Cambridge"}
PUBL_EN = {"Мысль": "Mysl’ Publ.", "Издательский центр «Академия»": "Akademiya Publ.",
           "Наука, Сибирское отделение": "Nauka Publ.", "Cambridge University Press": "CUP"}
PUBL_RU = {"Издательский центр «Академия»": "Академия", "Наука, Сибирское отделение": "Наука",
           "Cambridge University Press": "Cambridge Univ. Press"}
MAX_ALL_AUTHORS = 15  # journal rule: all authors, or the first three if more than 15


def authors_str(auth, etal, lang, sep=", "):
    """lang: 'ru' (Cyrillic as is), 'lat' (Latin as is or transliterated)."""
    lst = auth
    if len(lst) > MAX_ALL_AUTHORS:
        lst, etal = lst[:3], True
    parts = []
    for last, ini in lst:
        if lang == "lat" and is_cyrillic(last):
            last, ini = translit(last), translit_initials(ini)
        parts.append(f"{last} {ini}".strip())
    s = sep.join(parts)
    if etal:
        s += ", et al." if lang == "lat" or not is_cyrillic(lst[0][0]) else " и др."
    return s


def title_of(key, e):
    return clean(TITLE_OVERRIDE.get(key, e.get("title", "")))


def journal_of(e):
    j = clean(e.get("journaltitle") or e.get("journal") or "")
    return JOURNAL_ABBR.get(j, (j, j))


def doi_ru(e):
    return f" doi: {e['doi']}" if e.get("doi") else ""


def gost_entry(key, e):
    """Traditional list (GOST, journal examples). Authors in italics."""
    auth, etal = split_authors(e["author"])
    ru = is_cyrillic(auth[0][0])
    a = authors_str(auth, etal, "ru")
    t = title_of(key, e)
    typ = e["_type"]
    yr = e["year"]
    if typ == "article":
        jr = journal_of(e)[0]
        vol = e.get("volume")
        num = e.get("number")
        pg = clean(e.get("pages", ""))
        bits = [jr if jr.endswith(".") else jr + ".", f"{yr}."]
        if vol:
            bits.append(("Т. " if ru else "Vol. ") + vol + ".")
        if num:
            bits.append("№ " + num.replace("-", "–") + ".")
        if pg:
            bits.append(("С. " if ru else ("P. " if "–" in pg else "Art. ")) + pg + ".")
        return f"\\textit{{{a}}} {t} // " + " ".join(bits) + doi_ru(e)
    if typ == "book":
        city = CITY_RU.get(e.get("location", ""), e.get("location", ""))
        pub = PUBL_RU.get(e.get("publisher", ""), e.get("publisher", ""))
        pt = e.get("pagetotal")
        s = f"\\textit{{{a}}} {t}. {city}: {pub}, {yr}."
        if pt:
            s += f" {pt} {'с' if ru else 'p'}."
        return s + doi_ru(e)
    if typ == "incollection":
        city = CITY_RU.get(e.get("location", ""), e.get("location", ""))
        pub = e.get("publisher", "")
        return (f"\\textit{{{a}}} {t} // {clean(e['booktitle'])}. {city}: {pub}, {yr}. "
                f"С. {clean(e['pages'])}.")
    if typ == "online":
        return (f"\\textit{{{a}}} {t}. {e.get('eprinttype', 'arXiv')} preprint, {yr}. "
                f"arXiv:{e['eprint']}.{doi_ru(e)}")
    raise ValueError(typ)


def refs_entry(key, e):
    """Latin References list (journal rules)."""
    auth, etal = split_authors(e["author"])
    ru = is_cyrillic(auth[0][0])
    a = authors_str(auth, etal, "lat")
    typ = e["_type"]
    yr = e["year"]
    doi = f"\\\\ https://doi.org/{e['doi']}" if e.get("doi") else ""
    if typ == "article":
        t = EN_TITLE.get(key) if ru else title_of(key, e)
        jr = journal_of(e)[1]
        bits = [yr]
        if e.get("volume"):
            bits.append("vol. " + e["volume"])
        if e.get("number"):
            bits.append("no. " + e["number"].replace("-", "–"))
        pg = clean(e.get("pages", ""))
        if pg:
            bits.append(("pp. " if "–" in pg else "art. ") + pg)
        tt = t if t[-1] in "?!." else t + "."
        s = f"{a} {tt} \\textit{{{jr}}}, " + ", ".join(bits) + "."
        if ru:
            s += " (In Russ.)."
        return s + doi
    if typ == "book":
        city = CITY_EN.get(e.get("location", ""), e.get("location", ""))
        pub = PUBL_EN.get(e.get("publisher", ""), e.get("publisher", ""))
        if ru:
            t = f"\\textit{{{translit(clean(e['title']))}}} [{EN_TITLE[key]}]"
        else:
            t = f"\\textit{{{title_of(key, e)}}}"
        s = f"{a} {t}. {city}: {pub}, {yr}."
        if e.get("pagetotal"):
            s += f" {e['pagetotal']} p."
        return s + doi
    if typ == "incollection":
        city = CITY_EN.get(e.get("location", ""), e.get("location", ""))
        pub = PUBL_EN.get(e.get("publisher", ""), e.get("publisher", ""))
        bt, bt_en = EN_BOOKTITLE[key]
        return (f"{a} {EN_TITLE[key]}. In \\textit{{{bt}}} [{bt_en}]. {city}: {pub}, {yr}, "
                f"pp. {clean(e['pages'])}. (In Russ.).")
    if typ == "online":
        return (f"{a} {title_of(key, e)}. arXiv preprint, {yr}, arXiv:{e['eprint']}.{doi}")
    raise ValueError(typ)


def cite_label(key, e):
    auth, etal = split_authors(e["author"])
    ru = is_cyrillic(auth[0][0])
    n = len(auth) + (1 if etal else 0)
    if n == 1:
        names = auth[0][0]
    elif n == 2:
        names = f"{auth[0][0]}, {auth[1][0]}"
    else:
        names = auth[0][0] + (" и др." if ru else " et al.")
    return names, e["year"]


def sort_key_ru(label):
    return (0 if is_cyrillic(label) else 1, label.lower())


def render_cite(keys, entries):
    labs = [cite_label(k, entries[k]) for k in keys]
    labs.sort(key=lambda x: (sort_key_ru(x[0]), x[1]))
    out, i = [], 0
    while i < len(labs):
        name, yrs = labs[i][0], [labs[i][1]]
        while i + 1 < len(labs) and labs[i + 1][0] == name:
            i += 1
            yrs.append(labs[i][1])
        out.append(f"{name}, {', '.join(yrs)}")
        i += 1
    return "(" + "; ".join(out) + ")"


# ------------------------------------------------------------------------------ main
def build_tex():
    entries = parse_bib(BIB)
    src = SRC.read_text(encoding="utf-8")
    used: list[str] = []

    def repl(m):
        keys = [k.strip() for k in m.group(1).split(",")]
        for k in keys:
            if k not in entries:
                sys.exit(f"unknown key {k}")
            if k not in used:
                used.append(k)
        return render_cite(keys, entries)

    text = re.sub(r"\\cite\{([^}]*)\}", repl, src)

    def sort_gost(k):
        auth, _ = split_authors(entries[k]["author"])
        last = auth[0][0]
        return (0 if is_cyrillic(last) else 1, last.lower(), entries[k]["year"])

    def sort_refs(k):
        auth, _ = split_authors(entries[k]["author"])
        last = auth[0][0]
        return ((translit(last) if is_cyrillic(last) else last).lower(), entries[k]["year"])

    gost = "\n\n".join("\\noindent " + gost_entry(k, entries[k]) for k in sorted(used, key=sort_gost))
    refs = "\n\n".join("\\noindent " + refs_entry(k, entries[k]) for k in sorted(used, key=sort_refs))
    text = text.replace("%%GOST%%", gost).replace("%%REFS%%", refs)

    # journal typography checks on the body (outside the bibliographies)
    body = text.split("%%BODY-END%%")[0]
    problems = []
    for ch, what in (("«", "guillemets"), ("»", "guillemets"), ("ё", "yo"), ("Ё", "YO")):
        if ch in body:
            problems.append(f"{what}: {body.count(ch)}")
    dc = re.findall(r"\d\{,\}\d|\d,\d", body)
    if dc:
        problems.append(f"decimal commas: {len(dc)} e.g. {dc[:5]}")
    if problems:
        sys.exit("typography check failed: " + "; ".join(problems))
    OUT_TEX.write_text(text, encoding="utf-8")
    print(f"wrote {OUT_TEX.name}: {len(used)} sources cited")
    return text


def char_count(text: str):
    after = text.split(r"\begin{document}")[1]
    b = after.split("%%BODY-END%%")[0]
    tail = after.split("%%BODY-END%%")[1] if "%%BODY-END%%" in after else ""
    tabs = re.findall(r"\\begin\{tabular\}.*?\\end\{tabular\}", after, flags=re.S)
    tabs += re.findall(r"^\\noindent Таблица \d\..*$", tail, flags=re.M)
    figs = re.findall(r"^\\noindent Рис\. \d\..*$", tail, flags=re.M)

    def detex(x):
        x = re.sub(r"%.*", "", x)
        x = re.sub(r"\\(?:label|ref|url|includegraphics)\{[^}]*\}", "", x)
        x = re.sub(r"\\[a-zA-Z]+\*?(\[[^\]]*\])?", "", x)
        x = re.sub(r"[{}$~&]", "", x)
        return re.sub(r"\s+", " ", x).strip()

    t = b
    for x in tabs + figs:
        t = t.replace(x, "")
    # cut bibliographies / english block if any left inside
    n_text = len(detex(t))
    n_tab = sum(len(detex(x)) for x in tabs)
    print(f"characters: text {n_text}, tables {n_tab}, total {n_text + n_tab} (limit 45000)")


def build_docx():
    figs_dir = MAN / "figures"
    # run inside manuscript_v2 with a relative resource path: with an absolute
    # --resource-path pandoc 3.9 leaves a few \overline formulas as raw TeX
    cmd = ["pandoc", OUT_TEX.name, "-o", OUT_DOCX.name, "--from", "latex", "--to", "docx",
           "--resource-path=.", "--wrap=none"]
    subprocess.run(cmd, check=True, cwd=MAN)
    try:
        import docx  # type: ignore
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_LINE_SPACING
    except ImportError:
        print("python-docx not installed; docx left as pandoc produced it")
        return
    d = docx.Document(str(OUT_DOCX))
    for s in d.sections:
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)
    for st in d.styles:
        try:
            st.font.name = "Times New Roman"
            rpr = st.element.get_or_add_rPr()
            rf = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
            if rf is None:
                from docx.oxml.ns import qn
                from docx.oxml import OxmlElement
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            from docx.oxml.ns import qn
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), "Times New Roman")
        except Exception:
            pass
    body_styles = {"Normal", "Body Text", "First Paragraph", "Compact"}
    for st in d.styles:
        if st.name in body_styles:
            st.font.size = Pt(12)
            pf = st.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            pf.first_line_indent = Cm(1.25)
            pf.space_after = Pt(0)
    for p in d.paragraphs:
        for r in p.runs:
            r.font.name = "Times New Roman"
            if r.font.size is None or r.font.size > Pt(12):
                r.font.size = Pt(12)
    d.save(str(OUT_DOCX))
    print(f"wrote {OUT_DOCX.name}")


def build_pdf():
    if not shutil.which("latexmk"):
        print("latexmk not found; skipping PDF")
        return
    subprocess.run(["latexmk", "-xelatex", "-interaction=nonstopmode", "-halt-on-error",
                    "-quiet", OUT_TEX.name], cwd=MAN, check=True)
    subprocess.run(["latexmk", "-c", "-quiet", OUT_TEX.name], cwd=MAN, check=False)
    print("wrote article1_v7.pdf")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--no-docx", action="store_true")
    a = ap.parse_args()
    txt = build_tex()
    char_count(txt)
    if not a.no_pdf:
        build_pdf()
    if not a.no_docx:
        build_docx()
